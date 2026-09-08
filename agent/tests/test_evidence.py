from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from app.models.evidence import LiteratureSearchInput
from app.services.evidence_importer import COLUMN_MAP, SMILES_COLUMNS, import_literature
from app.services.evidence_store import EvidenceStore


def make_fixture(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "EN"
    headers = [*COLUMN_MAP.keys(), *SMILES_COLUMNS]
    sheet.append(headers)
    first: dict[str, Any] = {
        "Title": "Licorice supramolecular assembly",
        "Link/DOI": "10.1000/licorice",
        "Year": 2024,
        "Authors": "A. Author",
        "Natural Medicine / Herb Pair / Formula": "Licorice (Gancao)",
        "Organic Small Molecule(s)": "Glycyrrhizic acid",
        "Assembly Morphology": "nanofiber",
        "Assembly Phase": "liquid",
        "Solvent Type": "water",
        "pH": "7.0",
        "Temperature": "25 C",
        "Supramolecular Interaction Types": "hydrogen bonding",
        "Mechanism": "deterministic fixture statement",
        SMILES_COLUMNS[0]: "CCO",
    }
    second: dict[str, Any] = {
        "Year": 2023,
        "Natural Medicine / Herb Pair / Formula": "Astragalus",
        "Assembly Morphology": "particle",
    }
    for payload in (first, second):
        sheet.append([payload.get(header) for header in headers])
    workbook.save(path)


def test_import_quality_provenance_and_missing_values(tmp_path: Path) -> None:
    source = tmp_path / "fixture.xlsx"
    make_fixture(source)
    records, report = import_literature(source)
    assert report.total_rows == report.imported_rows == 2
    assert report.missing_title == 1
    assert report.missing_link_or_doi == 1
    assert report.missing_compounds == 1
    assert report.rows_with_smiles == 1
    assert records[0].source_file == source.name
    assert records[0].sheet == "EN"
    assert records[0].source_row == 2
    assert records[1].title is None


def test_index_hash_idempotency_and_changed_source_rebuild(tmp_path: Path) -> None:
    source = tmp_path / "fixture.xlsx"
    database = tmp_path / "evidence.db"
    make_fixture(source)
    store = EvidenceStore(source, database)
    first_hash = store.quality_report().source_sha256
    assert store.ensure_index() is False
    store.close()

    workbook = load_workbook(source)
    active_sheet = workbook.active
    assert active_sheet is not None
    active_sheet.cell(row=3, column=1, value="Changed title")
    workbook.save(source)
    changed = EvidenceStore(source, database)
    try:
        assert changed.quality_report().source_sha256 != first_hash
        assert changed.search(LiteratureSearchInput(query="Changed")).hits
    finally:
        changed.close()


def test_hybrid_recall_smiles_and_deterministic_sort(tmp_path: Path) -> None:
    source = tmp_path / "fixture.xlsx"
    make_fixture(source)
    store = EvidenceStore(source, tmp_path / "evidence.db")
    try:
        herb = store.search(LiteratureSearchInput(herbs=["甘草"], top_k=10))
        compound = store.search(LiteratureSearchInput(compounds=["Glycyrrhizic acid"], top_k=10))
        smiles = store.search(LiteratureSearchInput(smiles=["CCO"], top_k=10))
        repeated = store.search(LiteratureSearchInput(query="assembly", top_k=10))
        assert herb.hits[0].document_id == "EN:2"
        assert compound.hits[0].matched_fields == ["compounds"]
        assert smiles.hits[0].channel_scores["smiles"] == 1.0
        assert [hit.document_id for hit in repeated.hits] == [
            hit.document_id
            for hit in store.search(LiteratureSearchInput(query="assembly", top_k=10)).hits
        ]
        empty = store.search(LiteratureSearchInput(query="no-such-token"))
        assert empty.hits == []
        assert empty.degraded is True
    finally:
        store.close()
