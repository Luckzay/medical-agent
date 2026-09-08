from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Final

from openpyxl import load_workbook

from app.models.evidence import DataQualityReport, LiteratureRecord

SCHEMA_VERSION: Final = "1.0"

COLUMN_MAP: Final[dict[str, str]] = {
    "Title": "title",
    "Link/DOI": "link_or_doi",
    "Year": "year",
    "Country": "country",
    "Authors": "authors",
    "First Affiliation": "first_affiliation",
    "Natural Medicine / Herb Pair / Formula": "herbs",
    "Organic Small Molecule(s)": "compounds",
    "Metal Ion(s)": "metal_ions",
    "Organic Macromolecule(s)": "organic_macromolecules",
    "TCM Decoction Multi-component Supramolecule (Composition)": "decoction_composition",
    "Main Composition of Supramolecular Assembly": "assembly_composition",
    "Assembly Morphology": "assembly_morphology",
    "Assembly Phase": "assembly_phase",
    "Solvent Type": "solvent_type",
    "pH": "ph",
    "Temperature": "temperature",
    "Supramolecular Interaction Types": "interaction_types",
    "Intermolecular Interaction Sites": "interaction_sites",
    "Average Size": "average_size",
    "PDI": "pdi",
    "Zeta Potential": "zeta_potential",
    "CAC/CMC": "cac_cmc",
    "Tgel": "tgel",
    "Clinical Problem": "clinical_problem",
    "Pharmacodynamic Indicators": "pharmacodynamic_indicators",
    "Efficacy Comparison (Specific Indicators)": "efficacy_comparison",
    "Mechanism": "mechanism",
    "Remarks": "remarks",
}
SMILES_COLUMNS: Final = tuple(f"Small Molecule {number} SMILES (PubChem)" for number in range(1, 5))


def source_sha256(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    rendered = str(value).strip()
    return rendered or None


def _year(value: object) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = _text(value)
    if text is not None and text.isdigit():
        return int(text)
    return None


def import_literature(path: str | Path) -> tuple[list[LiteratureRecord], DataQualityReport]:
    source = Path(path)
    digest = source_sha256(source)
    workbook = load_workbook(source, read_only=True, data_only=True)
    records: list[LiteratureRecord] = []
    total_rows = 0
    coverage_counts = {field: 0 for field in COLUMN_MAP.values()}
    coverage_counts["small_molecule_smiles"] = 0
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            try:
                headers = [str(value).strip() if value is not None else "" for value in next(rows)]
            except StopIteration:
                continue
            positions = {header: index for index, header in enumerate(headers)}
            required = set(COLUMN_MAP) | set(SMILES_COLUMNS)
            missing = sorted(required - positions.keys())
            if missing:
                raise ValueError(
                    f"Excel sheet '{sheet.title}' missing columns: {', '.join(missing)}"
                )
            for source_row, values in enumerate(rows, start=2):
                if not any(_text(value) is not None for value in values):
                    continue
                total_rows += 1
                payload: dict[str, object] = {}
                for column, field in COLUMN_MAP.items():
                    value = values[positions[column]] if positions[column] < len(values) else None
                    normalized: object = _year(value) if field == "year" else _text(value)
                    payload[field] = normalized
                    if normalized is not None:
                        coverage_counts[field] += 1
                smiles = [
                    text
                    for column in SMILES_COLUMNS
                    if (text := _text(values[positions[column]])) is not None
                ]
                if smiles:
                    coverage_counts["small_molecule_smiles"] += 1
                records.append(
                    LiteratureRecord.model_validate(
                        {
                            "document_id": f"{sheet.title}:{source_row}",
                            "small_molecule_smiles": smiles,
                            "source_file": source.name,
                            "sheet": sheet.title,
                            "source_row": source_row,
                            **payload,
                        }
                    )
                )
    finally:
        workbook.close()
    imported = len(records)
    denominator = imported or 1
    report = DataQualityReport(
        total_rows=total_rows,
        imported_rows=imported,
        missing_title=sum(record.title is None for record in records),
        missing_link_or_doi=sum(record.link_or_doi is None for record in records),
        missing_herb=sum(record.herbs is None for record in records),
        missing_compounds=sum(record.compounds is None for record in records),
        rows_with_smiles=sum(bool(record.small_molecule_smiles) for record in records),
        field_coverage={
            field: round(count / denominator, 6) for field, count in coverage_counts.items()
        },
        source_sha256=digest,
        schema_version=SCHEMA_VERSION,
    )
    return records, report
